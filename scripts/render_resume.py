"""
Renderer script to generate tailored resume.docx and resume.pdf.
Uses python-docx and docx2pdf (via MS Word) to generate clean ATS-safe 1-page resumes.
"""

import os
import sys
import yaml
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx2pdf import convert

def add_bottom_border(paragraph, color_hex="1E293B", sz="8"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="2" w:color="{color_hex}"/></w:pBdr>')
    pPr.append(pBdr)

def load_source_of_truth():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    sot_dir = os.path.join(base_dir, "source_of_truth")
    
    with open(os.path.join(sot_dir, "profile.yaml"), "r", encoding="utf-8") as f:
        profile = yaml.safe_load(f)
    with open(os.path.join(sot_dir, "skills.yaml"), "r", encoding="utf-8") as f:
        skills = yaml.safe_load(f)
    with open(os.path.join(sot_dir, "bullets_bank.yaml"), "r", encoding="utf-8") as f:
        bullets_bank = yaml.safe_load(f)
    with open(os.path.join(sot_dir, "certs_projects.yaml"), "r", encoding="utf-8") as f:
        certs_proj = yaml.safe_load(f)

    bullets_by_id = {b["id"]: b for b in bullets_bank["bullets"]}
    return profile, skills, bullets_by_id, certs_proj

def render_data_engineer(config, out_docx):
    profile, _, bullets_by_id, certs_proj = load_source_of_truth()
    
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.38)
    section.bottom_margin = Inches(0.32)
    section.left_margin = Inches(0.50)
    section.right_margin = Inches(0.50)

    primary_color = RGBColor(15, 23, 42)
    text_color = RGBColor(30, 41, 59)
    muted_color = RGBColor(71, 85, 105)

    # Header Name
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    p_name.paragraph_format.line_spacing = 1.0
    r_name = p_name.add_run(profile.get("name", "HRUSHIKESH WADEKAR"))
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(18)
    r_name.font.bold = True
    r_name.font.color.rgb = primary_color

    # Header Contact
    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(7)
    p_contact.paragraph_format.line_spacing = 1.0
    c = profile["contact"]
    contact_text = f"{c.get('location', 'Mumbai, Maharashtra, India')}  |  {c['phone']}  |  {c['email']}  |  {c.get('linkedin', 'linkedin.com/in/hrushikeshwadekar')}"
    r_contact = p_contact.add_run(contact_text)
    r_contact.font.name = "Calibri"
    r_contact.font.size = Pt(8.5)
    r_contact.font.color.rgb = muted_color

    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.0
        add_bottom_border(p, color_hex="1E293B", sz="8")
        r = p.add_run(title)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = primary_color
        return p

    # 1. Summary
    add_section_header("PROFESSIONAL SUMMARY")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_before = Pt(1)
    p_sum.paragraph_format.space_after = Pt(3)
    p_sum.paragraph_format.line_spacing = 1.07
    
    summary_runs = config.get("summary_runs")
    if summary_runs:
        for item in summary_runs:
            text = item.get("text", "")
            bold = item.get("bold", False)
            r = p_sum.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            r.font.bold = bold
            r.font.color.rgb = text_color
    else:
        summary_text = config.get("summary", "")
        r = p_sum.add_run(summary_text)
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.color.rgb = text_color

    # 2. Certifications
    add_section_header("CERTIFICATIONS")
    for cert in certs_proj.get("certifications", []):
        p_c = doc.add_paragraph()
        p_c.paragraph_format.space_before = Pt(1)
        p_c.paragraph_format.space_after = Pt(0)
        p_c.paragraph_format.line_spacing = 1.02
        p_c.paragraph_format.left_indent = Inches(0.18)
        
        r_b = p_c.add_run("•  ")
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(8.5)
        r_b.font.bold = True
        
        r_name = p_c.add_run(cert["name"])
        r_name.font.name = "Calibri"
        r_name.font.size = Pt(8.5)
        r_name.font.bold = True
        
        r_sep = p_c.add_run(" — ")
        r_sep.font.name = "Calibri"
        r_sep.font.size = Pt(8.5)
        
        r_iss = p_c.add_run(cert["issuer"])
        r_iss.font.name = "Calibri"
        r_iss.font.size = Pt(8.5)
        r_iss.font.italic = True
        
        r_sp = p_c.add_run(f"    (Issued: {cert['issue_date']})")
        r_sp.font.name = "Calibri"
        r_sp.font.size = Pt(8.0)
        r_sp.font.color.rgb = muted_color

        p_id = doc.add_paragraph()
        p_id.paragraph_format.space_before = Pt(0)
        p_id.paragraph_format.space_after = Pt(1)
        p_id.paragraph_format.line_spacing = 1.0
        p_id.paragraph_format.left_indent = Inches(0.32)
        r_id = p_id.add_run(f"ID: {cert['credential_id']}")
        r_id.font.name = "Consolas"
        r_id.font.size = Pt(7.5)
        r_id.font.color.rgb = muted_color

    # 3. Technical Skills
    add_section_header("TECHNICAL SKILLS")
    skills_lines = config.get("skills_lines", [])
    for cat, items in skills_lines:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_before = Pt(0)
        p_s.paragraph_format.space_after = Pt(1)
        p_s.paragraph_format.line_spacing = 1.04
        
        r_cat = p_s.add_run(f"{cat}: ")
        r_cat.font.name = "Calibri"
        r_cat.font.size = Pt(8.5)
        r_cat.font.bold = True
        r_cat.font.color.rgb = primary_color
        
        r_it = p_s.add_run(items)
        r_it.font.name = "Calibri"
        r_it.font.size = Pt(8.5)
        r_it.font.color.rgb = text_color

    # 4. Work Experience
    add_section_header("WORK EXPERIENCE")
    p_exp_hdr = doc.add_paragraph()
    p_exp_hdr.paragraph_format.space_before = Pt(2)
    p_exp_hdr.paragraph_format.space_after = Pt(0)
    p_exp_hdr.paragraph_format.line_spacing = 1.0
    
    r_role = p_exp_hdr.add_run("Data Engineer")
    r_role.font.name = "Calibri"
    r_role.font.size = Pt(9.0)
    r_role.font.bold = True
    
    r_sep = p_exp_hdr.add_run(" — ")
    r_sep.font.name = "Calibri"
    r_sep.font.size = Pt(9.0)
    
    r_comp = p_exp_hdr.add_run("Tata Consultancy Services (Client: State Bank of India)")
    r_comp.font.name = "Calibri"
    r_comp.font.size = Pt(9.0)
    r_comp.font.italic = True
    
    r_dt = p_exp_hdr.add_run("                       Apr 2024 – Present")
    r_dt.font.name = "Calibri"
    r_dt.font.size = Pt(8.5)
    r_dt.font.bold = True
    r_dt.font.color.rgb = primary_color

    p_proj_sub = doc.add_paragraph()
    p_proj_sub.paragraph_format.space_before = Pt(0)
    p_proj_sub.paragraph_format.space_after = Pt(2)
    p_proj_sub.paragraph_format.line_spacing = 1.0
    
    r_p_sub = p_proj_sub.add_run("Project: Vendors Payment & Settlement System (VPS)")
    r_p_sub.font.name = "Calibri"
    r_p_sub.font.size = Pt(8.5)
    r_p_sub.font.bold = True
    
    r_loc = p_proj_sub.add_run("                                        Mumbai, India")
    r_loc.font.name = "Calibri"
    r_loc.font.size = Pt(8.0)
    r_loc.font.italic = True
    r_loc.font.color.rgb = muted_color

    for b_info in config.get("experience_bullets", []):
        lead = b_info.get("lead", "")
        text = b_info.get("text", "")
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_before = Pt(1)
        p_b.paragraph_format.space_after = Pt(1)
        p_b.paragraph_format.line_spacing = 1.05
        p_b.paragraph_format.left_indent = Inches(0.18)
        
        r_dot = p_b.add_run("•  ")
        r_dot.font.name = "Calibri"
        r_dot.font.size = Pt(8.5)
        
        if lead:
            r_ld = p_b.add_run(f"{lead}: ")
            r_ld.font.name = "Calibri"
            r_ld.font.size = Pt(8.5)
            r_ld.font.bold = True
        
        r_tx = p_b.add_run(text)
        r_tx.font.name = "Calibri"
        r_tx.font.size = Pt(8.5)
        r_tx.font.color.rgb = text_color

    # 5. Technical Projects
    add_section_header("TECHNICAL PROJECTS")
    p_pr_title = doc.add_paragraph()
    p_pr_title.paragraph_format.space_before = Pt(2)
    p_pr_title.paragraph_format.space_after = Pt(0)
    p_pr_title.paragraph_format.line_spacing = 1.0
    r_prt = p_pr_title.add_run("Automated REST API Batch Ingestion & Dimensional Warehouse Pipeline")
    r_prt.font.name = "Calibri"
    r_prt.font.size = Pt(8.8)
    r_prt.font.bold = True

    p_pr_tech = doc.add_paragraph()
    p_pr_tech.paragraph_format.space_before = Pt(0)
    p_pr_tech.paragraph_format.space_after = Pt(2)
    p_pr_tech.paragraph_format.line_spacing = 1.0
    r_tech = p_pr_tech.add_run("Technologies: Python, PySpark, Snowflake, AWS S3, SQL, Dimensional Modeling (Star Schema)")
    r_tech.font.name = "Calibri"
    r_tech.font.size = Pt(7.8)
    r_tech.font.italic = True
    r_tech.font.color.rgb = muted_color

    for b_info in config.get("project_bullets", []):
        text = b_info.get("text", "") if isinstance(b_info, dict) else str(b_info)
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_before = Pt(1)
        p_b.paragraph_format.space_after = Pt(1)
        p_b.paragraph_format.line_spacing = 1.05
        p_b.paragraph_format.left_indent = Inches(0.18)
        
        r_dot = p_b.add_run("•  ")
        r_dot.font.name = "Calibri"
        r_dot.font.size = Pt(8.5)
        
        r_tx = p_b.add_run(text)
        r_tx.font.name = "Calibri"
        r_tx.font.size = Pt(8.5)
        r_tx.font.color.rgb = text_color

    # 6. Education
    add_section_header("EDUCATION")
    p_edu1 = doc.add_paragraph()
    p_edu1.paragraph_format.space_before = Pt(1)
    p_edu1.paragraph_format.space_after = Pt(0)
    p_edu1.paragraph_format.line_spacing = 1.0
    r_e1_d = p_edu1.add_run("Master of Computer Applications (MCA)")
    r_e1_d.font.name = "Calibri"
    r_e1_d.font.size = Pt(8.5)
    r_e1_d.font.bold = True
    r_e1_dt = p_edu1.add_run("                                                      2025 – 2027 (Pursuing)")
    r_e1_dt.font.name = "Calibri"
    r_e1_dt.font.size = Pt(8.5)
    r_e1_dt.font.bold = True

    p_edu1_inst = doc.add_paragraph()
    p_edu1_inst.paragraph_format.space_before = Pt(0)
    p_edu1_inst.paragraph_format.space_after = Pt(2)
    p_edu1_inst.paragraph_format.line_spacing = 1.0
    r_e1_in = p_edu1_inst.add_run("G. H. Raisoni College of Engineering and Management (KBC NMU), Jalgaon")
    r_e1_in.font.name = "Calibri"
    r_e1_in.font.size = Pt(8.0)
    r_e1_in.font.color.rgb = muted_color

    p_edu2 = doc.add_paragraph()
    p_edu2.paragraph_format.space_before = Pt(1)
    p_edu2.paragraph_format.space_after = Pt(0)
    p_edu2.paragraph_format.line_spacing = 1.0
    r_e2_d = p_edu2.add_run("Bachelor of Computer Applications (BCA)")
    r_e2_d.font.name = "Calibri"
    r_e2_d.font.size = Pt(8.5)
    r_e2_d.font.bold = True
    r_e2_dt = p_edu2.add_run("                                                      2020 – 2023")
    r_e2_dt.font.name = "Calibri"
    r_e2_dt.font.size = Pt(8.5)
    r_e2_dt.font.bold = True

    p_edu2_inst = doc.add_paragraph()
    p_edu2_inst.paragraph_format.space_before = Pt(0)
    p_edu2_inst.paragraph_format.space_after = Pt(2)
    p_edu2_inst.paragraph_format.line_spacing = 1.0
    r_e2_in = p_edu2_inst.add_run("R. C. Patel Arts, Commerce and Science College (KBC NMU), Shirpur  |  CGPA: 9.62 / 10.0")
    r_e2_in.font.name = "Calibri"
    r_e2_in.font.size = Pt(8.0)
    r_e2_in.font.color.rgb = muted_color

    os.makedirs(os.path.dirname(out_docx), exist_ok=True)
    doc.save(out_docx)

def render_oracle_dba(config, out_docx):
    profile, _, bullets_by_id, _ = load_source_of_truth()
    
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.40)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    primary_color = RGBColor(0, 0, 0)
    text_color = RGBColor(0, 0, 0)

    # Name
    p_name = doc.add_paragraph()
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    p_name.paragraph_format.line_spacing = 1.0
    r_name = p_name.add_run(profile.get("name", "HRUSHIKESH WADEKAR"))
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(17)
    r_name.font.bold = True
    r_name.font.color.rgb = primary_color

    # Contact line
    p_contact = doc.add_paragraph()
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(6)
    p_contact.paragraph_format.line_spacing = 1.0
    c = profile["contact"]
    contact_text = f"{c['phone']} | {c['email']} | {c.get('location', 'Mumbai, Maharashtra, India')} | {c.get('linkedin', 'linkedin.com/in/hrushikeshwadekar')}"
    r_contact = p_contact.add_run(contact_text)
    r_contact.font.name = "Calibri"
    r_contact.font.size = Pt(9.0)
    r_contact.font.color.rgb = text_color

    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(title)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = primary_color
        return p

    # 1. Summary
    add_section_header("PROFESSIONAL SUMMARY")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_before = Pt(1)
    p_sum.paragraph_format.space_after = Pt(3)
    p_sum.paragraph_format.line_spacing = 1.06
    r_sum = p_sum.add_run(config.get("summary", ""))
    r_sum.font.name = "Calibri"
    r_sum.font.size = Pt(9.0)

    # 2. Professional Experience
    add_section_header("PROFESSIONAL EXPERIENCE")
    p_job = doc.add_paragraph()
    p_job.paragraph_format.space_before = Pt(2)
    p_job.paragraph_format.space_after = Pt(0)
    p_job.paragraph_format.line_spacing = 1.0
    r_j = p_job.add_run("Tata Consultancy Services (TCS) | Database Administrator")
    r_j.font.name = "Calibri"
    r_j.font.size = Pt(9.5)
    r_j.font.bold = True

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(2)
    p_sub.paragraph_format.line_spacing = 1.0
    r_sub = p_sub.add_run("Client: State Bank of India | Apr 2024 – Present")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(9.0)

    for b in config.get("experience_bullets", []):
        text = b.get("text", "") if isinstance(b, dict) else str(b)
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_before = Pt(0.5)
        p_b.paragraph_format.space_after = Pt(0.5)
        p_b.paragraph_format.line_spacing = 1.04
        p_b.paragraph_format.left_indent = Inches(0.18)
        
        r_dot = p_b.add_run("•  ")
        r_dot.font.name = "Calibri"
        r_dot.font.size = Pt(9.0)
        
        r_t = p_b.add_run(text)
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(9.0)

    # 3. Technical Skills
    add_section_header("TECHNICAL SKILLS")
    for cat, text in config.get("skills_lines", []):
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_before = Pt(0.5)
        p_s.paragraph_format.space_after = Pt(0.5)
        p_s.paragraph_format.line_spacing = 1.04
        p_s.paragraph_format.left_indent = Inches(0.18)
        
        r_dot = p_s.add_run("•  ")
        r_dot.font.name = "Calibri"
        r_dot.font.size = Pt(9.0)
        
        r_cat = p_s.add_run(f"{cat}: ")
        r_cat.font.name = "Calibri"
        r_cat.font.size = Pt(9.0)
        r_cat.font.bold = True
        
        r_val = p_s.add_run(text)
        r_val.font.name = "Calibri"
        r_val.font.size = Pt(9.0)

    # 4. Education
    add_section_header("EDUCATION")
    
    # MCA
    p_edu1 = doc.add_paragraph()
    p_edu1.paragraph_format.space_before = Pt(1)
    p_edu1.paragraph_format.space_after = Pt(0)
    p_edu1.paragraph_format.line_spacing = 1.0
    r_e1_d = p_edu1.add_run("Master of Computer Applications (MCA) — ")
    r_e1_d.font.name = "Calibri"
    r_e1_d.font.size = Pt(9.0)
    r_e1_d.font.bold = True
    r_e1_i = p_edu1.add_run("G. H. Raisoni College of Engineering and Management (KBC NMU), Jalgaon")
    r_e1_i.font.name = "Calibri"
    r_e1_i.font.size = Pt(9.0)
    r_e1_dt = p_edu1.add_run("  (2025 – 2027 Pursuing)")
    r_e1_dt.font.name = "Calibri"
    r_e1_dt.font.size = Pt(8.5)
    r_e1_dt.font.italic = True

    # BCA
    p_edu2 = doc.add_paragraph()
    p_edu2.paragraph_format.space_before = Pt(0.5)
    p_edu2.paragraph_format.space_after = Pt(1)
    p_edu2.paragraph_format.line_spacing = 1.0
    r_e2_d = p_edu2.add_run("Bachelor of Computer Applications (BCA) — ")
    r_e2_d.font.name = "Calibri"
    r_e2_d.font.size = Pt(9.0)
    r_e2_d.font.bold = True
    r_e2_i = p_edu2.add_run("R. C. Patel Arts, Commerce and Science College (KBC NMU), Shirpur")
    r_e2_i.font.name = "Calibri"
    r_e2_i.font.size = Pt(9.0)
    r_e2_dt = p_edu2.add_run("  (2020 – 2023 | CGPA: 9.62 / 10.0)")
    r_e2_dt.font.name = "Calibri"
    r_e2_dt.font.size = Pt(8.5)
    r_e2_dt.font.italic = True

    os.makedirs(os.path.dirname(out_docx), exist_ok=True)
    doc.save(out_docx)

def render(config_path, out_dir):
    with open(config_path, "r", encoding="utf-8") as f:
        config = yaml.safe_load(f)

    os.makedirs(out_dir, exist_ok=True)
    resume_type = config.get("resume_type", "data_engineer")
    out_docx = os.path.join(out_dir, "resume.docx")
    out_pdf = os.path.join(out_dir, "resume.pdf")

    if resume_type == "data_engineer":
        render_data_engineer(config, out_docx)
    elif resume_type == "oracle_dba":
        render_oracle_dba(config, out_docx)
    else:
        raise ValueError(f"Unknown resume type: {resume_type}")

    print(f"Generated DOCX: {out_docx}")
    convert(out_docx, out_pdf)
    print(f"Generated PDF: {out_pdf}")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python render_resume.py <tailoring_config.yaml> <output_dir>")
        sys.exit(1)
    render(sys.argv[1], sys.argv[2])
